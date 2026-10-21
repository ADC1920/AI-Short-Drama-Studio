# -*- coding: utf-8 -*-
"""md2docx 代码块视觉后处理（article-format 生成链固定后处理）。

md2docx 的 ``` 围栏只换等宽字体（Consolas），没有底纹和边框，提示词这种
要整段复制的内容与正文分不清。本脚本对代码块段落补三样：

1. 浅底纹 w:shd（默认 F6F3EC）+ 左侧竖线 w:pBdr（默认 C8BFA8）；
2. 左缩进 0.3cm、首行缩进清零（复制时行首不带空格）；
3. w:keepLines（段落不跨页拆断，防代码块被页边界切成两半）。

pPr 有 schema 顺序：pBdr/shd 必须插在 spacing/ind/jc 之前，keepLines 插在
pageBreakBefore 之前；本脚本用 insert_element_before 按序插入。

用法：python style_code_blocks.py <docx路径> [--dry-run] [--shd F6F3EC] [--border C8BFA8]
     python style_code_blocks.py --selftest
"""
import sys
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

PPR_AFTER_KEEP_LINES = (
    "w:pageBreakBefore", "w:framePr", "w:widowControl", "w:numPr",
    "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs",
    "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)
PPR_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)
PPR_AFTER_SHD = PPR_AFTER_PBDR[1:]

CODE_FONT = "Consolas"


def _is_code_run(run):
    rpr = run._element.rPr
    if rpr is None:
        return False
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        return False
    return fonts.get(qn("w:ascii")) == CODE_FONT


def _is_code_paragraph(par):
    runs = [r for r in par.runs if r.text.strip()]
    return bool(runs) and all(_is_code_run(r) for r in runs)


def _get_or_add(pPr, tag, after_tags):
    el = pPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        pPr.insert_element_before(el, *after_tags)
    return el


def style_code_blocks(path, shd_fill="F6F3EC", border_color="C8BFA8",
                      dry_run=False):
    doc = Document(path)
    count = 0
    for par in doc.paragraphs:
        if not _is_code_paragraph(par):
            continue
        count += 1
        if dry_run:
            continue
        pPr = par._p.get_or_add_pPr()

        _get_or_add(pPr, "w:keepLines", PPR_AFTER_KEEP_LINES)

        pBdr = pPr.find(qn("w:pBdr"))
        if pBdr is None:
            pBdr = OxmlElement("w:pBdr")
            pPr.insert_element_before(pBdr, *PPR_AFTER_PBDR)
        if pBdr.find(qn("w:left")) is None:
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "12")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), border_color)
            pBdr.append(left)

        shd = pPr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            pPr.insert_element_before(shd, *PPR_AFTER_SHD)
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), shd_fill)

        par.paragraph_format.left_indent = Cm(0.3)
        par.paragraph_format.first_line_indent = Cm(0)
    if not dry_run:
        doc.save(path)
    return count


def _selftest():
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("提示词代码块第一行")
    r.font.name = CODE_FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal = doc.add_paragraph("正文段落，不应被处理。")
    tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    tmp.close()
    doc.save(tmp.name)

    n = style_code_blocks(tmp.name)
    assert n == 1, f"应识别 1 个代码段，实际 {n}"

    d2 = Document(tmp.name)
    code_p = d2.paragraphs[0]
    pPr = code_p._p.pPr
    assert pPr.find(qn("w:shd")) is not None
    assert pPr.find(qn("w:pBdr")) is not None
    assert pPr.find(qn("w:keepLines")) is not None
    assert round(code_p.paragraph_format.left_indent.cm, 2) == 0.3
    # 正文段未被改动
    assert d2.paragraphs[1]._p.pPr is None or d2.paragraphs[1]._p.pPr.find(qn("w:shd")) is None

    # pPr 顺序：keepLines 必须早于 pBdr/shd/ind
    tags = [child.tag.split("}")[1] for child in pPr]
    assert tags.index("keepLines") < tags.index("pBdr") < tags.index("shd")
    assert tags.index("shd") < tags.index("ind")

    assert style_code_blocks(tmp.name, dry_run=True) == 1
    print("selftest PASS: 识别/样式/顺序/正文未改动/干跑 5 项")


if __name__ == "__main__":
    args = sys.argv[1:]
    if "--selftest" in args:
        _selftest()
        sys.exit(0)
    if not args or args[0].startswith("-"):
        print(__doc__)
        sys.exit(1)
    path = args[0]
    dry = "--dry-run" in args
    fill = args[args.index("--shd") + 1] if "--shd" in args else "F6F3EC"
    color = args[args.index("--border") + 1] if "--border" in args else "C8BFA8"
    n = style_code_blocks(path, fill, color, dry_run=dry)
    print(f"code_blocks={n} dry_run={dry} -> {path}")
